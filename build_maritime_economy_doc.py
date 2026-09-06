from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "Materi_Inti_Maritime_Economy_Sub-CPMK-4.docx"

BLUE = "1F4D78"
MID_BLUE = "2E74B5"
INK = "1F2937"
MUTED = "5B6472"


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Bullet Custom")
    p.add_run(text)
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph(style="Body")
    set_font(p.add_run(label), size=10.5, color=INK, bold=True)
    set_font(p.add_run(text), size=10.5, color=INK)
    return p


def add_sources(doc, sources):
    p = doc.add_paragraph(style="Sources Heading")
    set_font(p.add_run("Sitasi"), size=10.5, color=BLUE, bold=True)
    for source in sources:
        p = doc.add_paragraph(style="Source List")
        set_font(p.add_run(source), size=8.5, color=MUTED)


def add_section(doc, item):
    h = doc.add_paragraph(style="Heading 1")
    set_font(h.add_run(f"{item['no']}. {item['title']}"), size=16, color=MID_BLUE, bold=True)
    add_label_paragraph(doc, "Definisi singkat. ", item["definition"])
    p = doc.add_paragraph(style="Key Heading")
    set_font(p.add_run("Poin kunci"), size=10.5, color=BLUE, bold=True)
    for key in item["keys"]:
        add_bullet(doc, key)
    add_label_paragraph(doc, "Data/statistik. ", item["data"])
    add_label_paragraph(doc, "Contoh kasus Indonesia. ", item["case"])
    add_sources(doc, item["sources"])


sections = [
    {
        "no": 1,
        "title": "Definisi & Ruang Lingkup Maritime Economy",
        "definition": "Maritime economy adalah seluruh kegiatan ekonomi yang bergantung pada laut, pesisir, dan konektivitas antarpulau. Cakupannya lebih luas daripada perikanan: pemanfaatan sumber daya laut, jasa transportasi, industri pendukung, dan jasa lingkungan secara berkelanjutan.",
        "keys": [
            "Sektor utama meliputi perikanan, budidaya, pelayaran, pelabuhan, galangan kapal, wisata bahari, energi laut, dan bioteknologi kelautan.",
            "Nilai ekonomi dibentuk sepanjang rantai pasok: produksi, pengolahan, logistik, hingga pemasaran.",
            "Laut adalah ruang ekonomi sekaligus ruang ekologis; pemanfaatan harus mengikuti daya dukung ekosistem.",
            "UNCLOS memberi kerangka hak dan kewajiban negara atas ruang serta sumber daya laut; UU Kelautan menjadi kerangka nasional.",
        ],
        "data": "Kontribusi PDB perikanan pada triwulan IV 2024 tercatat 2,59%. Angka ini adalah indikator subsektor perikanan, bukan ukuran keseluruhan ekonomi maritim.",
        "case": "Pengelolaan WPPNRI menghubungkan tata ruang, pengelolaan penangkapan ikan, pelabuhan perikanan, dan rantai nilai hasil laut.",
        "sources": [
            "Republik Indonesia, UU No. 32 Tahun 2014 tentang Kelautan, 2014. https://peraturan.bpk.go.id/Details/38710/uu-no-32tahun-2014",
            "United Nations, United Nations Convention on the Law of the Sea (UNCLOS), 1982. https://www.un.org/depts/los/convention_agreements/texts/unclos/UNCLOS-TOC.htm",
            "Kementerian Kelautan dan Perikanan, Laporan Kinerja Pusdatin Triwulan I 2025 (memuat PDB perikanan triwulan IV 2024), 2025. https://www.kkp.go.id/storage/AkuntabilitasKinerja/3116/document-akuntabilitas-kinerja-68623a069e8b8LKJ%20Pusdatin%20TW%20I%202025.pdf",
        ],
    },
    {
        "no": 2,
        "title": "Maritime Civilization",
        "definition": "Peradaban maritim adalah pola kehidupan masyarakat yang membangun pengetahuan, teknologi, institusi, dan identitas melalui laut. Secara ekonomi, peradaban ini tumbuh ketika pelayaran, perdagangan, pelabuhan, dan penguasaan jalur laut menciptakan pertukaran barang serta budaya.",
        "keys": [
            "Laut berfungsi sebagai penghubung antarpulau dan jalur perdagangan internasional.",
            "Kemampuan navigasi, pembuatan kapal, dan pengamanan pelayaran menentukan daya saing ekonomi.",
            "Pelabuhan menjadi simpul perdagangan, pengumpulan komoditas, dan pertukaran budaya.",
            "Warisan maritim relevan bagi ekonomi modern: konektivitas, industri kapal, dan diplomasi maritim.",
        ],
        "data": "Bukti historis penting adalah Prasasti Kedukan Bukit bertarikh 682 M, yang sering digunakan dalam kajian awal Sriwijaya dan aktivitas baharinya.",
        "case": "Sriwijaya berkembang sebagai kekuatan maritim melalui jaringan perdagangan sekitar Selat Malaka; pelabuhan berfungsi sebagai entrepot bagi komoditas dari pedalaman dan perdagangan lintas kawasan.",
        "sources": [
            "Direktorat Pelindungan Kebudayaan, Mataram Kuna: Agraris atau Maritim, 2017. https://kebudayaan.kemdikbud.go.id/dpk/mataram-kuna-agraris-atau-maritim/",
            "Direktorat Pelindungan Kebudayaan, Cakrawala Mandala Dwipantara: Wawasan Kemaritiman Kerajaan Singhasari, 2017. https://kebudayaan.kemdikbud.go.id/dpk/cakrawala-ma%E1%B9%87%E1%B8%8Dala-dwipantara-wawasan-kemaritiman-kerajaan-singhasari/",
            "Direktorat Pelindungan Kebudayaan, Damar dalam Jaringan Perdagangan Masa Kerajaan Sriwijaya, 2017. https://kebudayaan.kemdikbud.go.id/dpk/damar-dalam-jaringan-perdagangan-masa-kerajaan-sriwijaya/",
        ],
    },
    {
        "no": 3,
        "title": "Human Resources Kemaritiman",
        "definition": "SDM kemaritiman mencakup nelayan, pembudidaya, awak kapal, pelaut, pekerja pelabuhan, pengolah hasil laut, ahli konservasi, dan tenaga industri perkapalan. Kualitas SDM menentukan produktivitas, keselamatan kerja, kepatuhan lingkungan, dan nilai tambah sektor maritim.",
        "keys": [
            "Kompetensi teknis perlu disertai sertifikasi keselamatan, navigasi, mutu, dan ketertelusuran produk.",
            "Pendidikan vokasi, pelatihan, penyuluhan, serta literasi digital dan keuangan memperkuat usaha pesisir.",
            "Pekerjaan layak menuntut perlindungan awak kapal, kontrak kerja, pengupahan, dan pencegahan kerja paksa.",
            "Kebutuhan keterampilan meningkat pada pelayaran, logistik maritim, galangan kapal, dan budidaya berbasis teknologi.",
        ],
        "data": "Survei ILO-BRIN 2024 menjangkau 3.396 pekerja kapal perikanan di 18 pelabuhan. Pelatihan KKP diikuti 39.057 peserta pada 2024.",
        "case": "Pendidikan vokasi dan sertifikasi KKP bagi anak nelayan, pembudidaya, serta pelaku pengolahan diarahkan untuk memperbesar akses mereka ke dunia usaha dan industri.",
        "sources": [
            "ILO dan BRIN, Memahami Kondisi Kerja Pekerja Kapal Penangkap Ikan di Indonesia, 2025. https://www.ilo.org/id/publications/memahami-kondisi-kerja-pekerja-kapal-penangkap-ikan-di-indonesia-bukti-dari",
            "ILO, Stocktaking Study for the Development of a Sectoral Skills Strategy: Indonesia Seafaring, 2023. https://www.ilo.org/publications/stocktaking-study-development-sectoral-skills-strategy-indonesia-seafaring",
            "KKP, KKP Maksimalkan Pendidikan Tinggi Vokasi Tingkatkan Kompetensi SDM Perikanan, 2025. https://kkp.go.id/news/news-detail/kkp-maksimalkan-pendidikan-tinggi-vokasi-tingkatkan-kompetensi-sdm-perikanan-z6PZ.html",
        ],
    },
    {
        "no": 4,
        "title": "Coastal Communities",
        "definition": "Masyarakat pesisir adalah komunitas yang kehidupan sosial-ekonominya terkait erat dengan sumber daya pesisir dan laut. Mata pencahariannya meliputi penangkapan ikan, budidaya, pengolahan, perdagangan, garam, wisata, dan jasa transportasi lokal.",
        "keys": [
            "Ekonomi pesisir umumnya didominasi usaha mikro-kecil dan rentan terhadap fluktuasi musim serta harga.",
            "Akses terhadap ruang pesisir, modal, pasar, teknologi, dan infrastruktur menentukan kesejahteraan.",
            "Kerusakan mangrove, terumbu karang, pencemaran, dan perubahan iklim langsung memengaruhi pendapatan.",
            "Tata kelola perlu memastikan partisipasi masyarakat, keadilan akses, dan diversifikasi mata pencaharian.",
        ],
        "data": "Laporan Bank Dunia memperkirakan sektor perikanan mendukung lebih dari 7 juta pekerjaan, memasok sekitar 50% protein nasional, dan bernilai sekitar US$26,9 miliar per tahun.",
        "case": "Di WPP 713, 714, dan 718, ketergantungan tinggi pada perikanan serta kapasitas adaptasi yang terbatas menjadikan masyarakat pesisir lebih rentan terhadap perubahan iklim.",
        "sources": [
            "Republik Indonesia, UU No. 27 Tahun 2007 tentang Pengelolaan Wilayah Pesisir dan Pulau-Pulau Kecil, 2007, jo. UU No. 1 Tahun 2014. https://peraturan.bpk.go.id/Details/39911/uu-no-27-tahun-2007 dan https://peraturan.bpk.go.id/Details/38521/uu-no-1-tahun-2014",
            "Kaczan, D. dkk., Hot Water Rising: The Impact of Climate Change on Indonesia's Fisheries and Coastal Communities, 2023. https://doi.org/10.1596/40564",
            "KKP, Ekonomi Biru Butuh UMKM yang Melek Keuangan, 2026 (memuat data KUSUKA 2024). https://www.kkp.go.id/djpdskp/ekonomi-biru-butuh-umkm-yang-melek-keuangan-GM8L/detail.html",
        ],
    },
    {
        "no": 5,
        "title": "Fisheries & Aquaculture",
        "definition": "Perikanan tangkap memanen sumber daya ikan dari perairan alami, sedangkan akuakultur membudidayakan organisme air dalam lingkungan terkelola. Keduanya berperan untuk pangan, lapangan kerja, ekspor, dan ekonomi pesisir, tetapi memerlukan pengelolaan stok serta lingkungan yang ketat.",
        "keys": [
            "Perikanan tangkap harus berbasis stok ikan, ukuran kapal/alat tangkap, musim, dan pengawasan.",
            "Budidaya dapat menambah pasokan, tetapi perlu menjaga kualitas air, benih, pakan, penyakit, dan limbah.",
            "Hilirisasi melalui rantai dingin, pengolahan, dan sertifikasi mutu meningkatkan nilai tambah.",
            "Keberlanjutan menuntut pencegahan IUU fishing, perlindungan habitat, dan ketertelusuran produk.",
        ],
        "data": "Produksi ikan budidaya 2024 mencapai 6,37 juta ton, naik 13,64% dari tahun sebelumnya; produksi rumput laut mencapai 10,80 juta ton.",
        "case": "Model budidaya rumput laut di Wakatobi seluas 50 ha menghasilkan nilai produksi sementara Rp1,09 miliar; ini menunjukkan peluang diversifikasi ekonomi pesisir.",
        "sources": [
            "KKP, Menteri Trenggono Berhasil Tingkatkan Produksi Perikanan Budi Daya 13,6% di 2024, 2024. https://kkp.go.id/news/news-detail/menteri-trenggono-berhasil-tingkatkan-produksi-perikanan-budi-daya-136-di-2024-vQq0.html",
            "BPS, Volume Produksi dan Nilai Produksi Perikanan Tangkap Menurut Provinsi dan Jenis Penangkapan, 2024, 2025. https://www.bps.go.id/id/statistics-table/3/U2k4d1MwcFZjRGhSVVRKaWRHRm5Temw1VURJeFp6MDkjMw%3D%3D/volume-produksi-dan-nilai-produksi-perikanan-tangkap-menurut-provinsi-dan-jenis-penangkapan--2024.html",
            "KKP, Kelautan dan Perikanan dalam Angka Tahun 2024, 2024. https://perpustakaan.kkp.go.id/knowledgerepository/index.php?id=1074693&p=show_detail",
        ],
    },
    {
        "no": 6,
        "title": "Maritime Tourism",
        "definition": "Wisata bahari adalah kegiatan perjalanan berbasis ekosistem dan lanskap laut, seperti selam, snorkeling, pesiar, pantai, pulau kecil, dan wisata budaya pesisir. Nilai ekonominya berasal dari jasa wisata, akomodasi, transportasi lokal, pemandu, UMKM, dan konservasi.",
        "keys": [
            "Daya tarik utama adalah terumbu karang, mangrove, lamun, pantai, biota laut, dan budaya pesisir.",
            "Konservasi adalah modal ekonomi: ekosistem yang rusak menurunkan kualitas destinasi.",
            "Pengelolaan perlu zonasi, daya dukung, aturan perilaku wisatawan, pengelolaan sampah, dan manfaat bagi warga.",
            "Wisata bahari berisiko menimbulkan overtourism bila kapasitas kunjungan tidak dikendalikan.",
        ],
        "data": "Data nasional 2024 mencatat 13,90 juta wisatawan mancanegara dan 1,021 miliar perjalanan wisatawan domestik; angka ini mencakup seluruh jenis wisata, bukan hanya wisata bahari.",
        "case": "Raja Ampat menghubungkan wisata selam dengan konservasi. Tutupan karang hidup di lokasi pemantauan meningkat dari 42,44% pada 2021 menjadi 48,29% pada 2024, disertai pembatasan aktivitas di Wayag.",
        "sources": [
            "Kementerian Pariwisata, SISPARNAS: Infografis Nasional Kepariwisataan, data 2024. https://sisparnas.kemenparekraf.go.id/p/54266",
            "KKP, KKP Berhasil Lindungi Populasi Pari & Hiu, Wisata Bahari Raja Ampat Makin Mendunia, 2025. https://www.kkp.go.id/news/news-detail/kkp-berhasil-lindungi-populasi-pari-hiu-wisata-bahari-raja-ampat-makin-mendunia.html",
            "KKP, Permen KP No. 26 Tahun 2024 tentang Kategori Kawasan Konservasi untuk Pariwisata Alam Perairan, 2024. https://jdih.kkp.go.id/Homedev/DetailPeraturan/6702",
        ],
    },
    {
        "no": 7,
        "title": "Shipping, Logistics & Shipbuilding",
        "definition": "Pelayaran menghubungkan pulau dan pasar; logistik mengelola arus barang, informasi, serta pergudangan; sedangkan galangan kapal membangun dan memperbaiki armada. Ketiganya menentukan biaya distribusi, keterhubungan wilayah, dan daya saing ekonomi kepulauan.",
        "keys": [
            "Pelabuhan adalah simpul integrasi laut-darat dan penentu kelancaran rantai pasok.",
            "Konektivitas memerlukan jadwal andal, muatan balik, hinterland, dan layanan logistik; tidak cukup hanya membuka rute.",
            "Galangan kapal mendukung kemandirian armada, pemeliharaan, keselamatan, dan penyerapan tenaga kerja terampil.",
            "Digitalisasi dokumen, keselamatan pelayaran, dan pengurangan emisi menjadi agenda modernisasi.",
        ],
        "data": "Pada 2024 terdapat 7 pelabuhan utama yang memenuhi standar dan 37 rute subsidi Tol Laut. Rute pelayaran yang saling terhubung tercapai 27%.",
        "case": "Program Tol Laut melayani wilayah terpencil melalui trayek tetap dan bersubsidi untuk memperbaiki distribusi barang serta mengurangi kesenjangan harga antarwilayah.",
        "sources": [
            "Republik Indonesia, UU No. 17 Tahun 2008 tentang Pelayaran, 2008. https://jdih.dephub.go.id/peraturan/detail?data=BIRXgzynMVj5Zo007DzuIi4KFw5WlEP8T4vQiLgyoCig8gi0Lz6n1We4jpLh4uJ4Z04uR56wikGnt4ZA1yWRpD5m49dA1sp8fcU49Y1ZSKk4dd7jmBrboqU2gKryjKzZoTo11Z6C3xK0RY4FMGACPCn5f5",
            "Kementerian Perhubungan, Laporan Kinerja Kementerian Perhubungan Tahun 2024, 2025. https://ppid.dephub.go.id/fileupload/informasi-berkala/20250516132623.LKIP_Kemenhub_2024.pdf",
            "BPS, Statistik Transportasi Laut 2024, 2025. https://www.bps.go.id/id/publication/2025/12/01/fdc6de7c2b34aa9b109edcad/statistik-transportasi-laut-2024.html",
        ],
    },
    {
        "no": 8,
        "title": "Blue Economy Indonesia",
        "definition": "Ekonomi biru adalah pendekatan pembangunan yang meningkatkan kesejahteraan dari sumber daya laut tanpa mengurangi kesehatan dan daya pulih ekosistem. Fokusnya bukan hanya pertumbuhan ekonomi, melainkan keseimbangan manfaat ekonomi, sosial, dan lingkungan.",
        "keys": [
            "Prioritas mencakup perikanan berkelanjutan, budidaya rendah dampak, konservasi, wisata bahari, energi laut, dan pengurangan sampah laut.",
            "Roadmap membutuhkan tata kelola lintas sektor, pembiayaan, data, inovasi, serta partisipasi masyarakat pesisir.",
            "Target keberhasilan diukur melalui PDB maritim, pekerjaan maritim, dan luas kawasan konservasi laut.",
            "Tantangan utama: perubahan iklim, IUU fishing, pencemaran plastik, konflik pemanfaatan ruang, dan ketimpangan manfaat.",
        ],
        "data": "Peta Jalan Ekonomi Biru menargetkan kontribusi PDB maritim naik dari 7,6% (kondisi awal) menjadi 15% pada 2045, seiring target kontribusi lapangan kerja maritim naik menjadi 12%. PDB perikanan triwulan IV 2024 adalah 2,59%, bukan ukuran keseluruhan ekonomi biru.",
        "case": "Peta Jalan Ekonomi Biru Indonesia membagi implementasi 2023-2045 ke lima fase: konsolidasi ekosistem, sumber pertumbuhan baru, diversifikasi, peningkatan daya saing, dan ekonomi biru berkelanjutan.",
        "sources": [
            "Bappenas, Peta Jalan Ekonomi Biru Indonesia Edisi 2, 2024. https://www.bappenas.go.id/index.php/unit-kerja/0405",
            "Bappenas, Indonesia Blue Economy Roadmap, 2023. https://perpustakaan.bappenas.go.id/e-library/file_upload/koleksi/migrasi-data-publikasi/file/Policy_Paper/ENG_Indonesia%20Blue%20Economy%20Roadmap_ebook.pdf",
            "Wuwung, L., McIlgorm, A., & Voyer, M., Sustainable ocean development policies in Indonesia: paving the pathways towards a maritime destiny, Frontiers in Marine Science, 2024 (mengutip target PDB maritim Peta Jalan Ekonomi Biru: 7,6% menjadi 15% pada 2045). https://doi.org/10.3389/fmars.2024.1401332",
            "Republik Indonesia, UU No. 32 Tahun 2014 tentang Kelautan, 2014. https://peraturan.bpk.go.id/Details/38710/uu-no-32tahun-2014",
        ],
    },
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.85)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

for style_name in ["Heading 1", "Heading 2"]:
    st = styles[style_name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

body = styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
body.base_style = styles["Normal"]
body.paragraph_format.space_after = Pt(5)
body.paragraph_format.line_spacing = 1.15

bullet = styles.add_style("Bullet Custom", WD_STYLE_TYPE.PARAGRAPH)
bullet.base_style = styles["List Bullet"]
bullet.paragraph_format.left_indent = Inches(0.25)
bullet.paragraph_format.first_line_indent = Inches(-0.15)
bullet.paragraph_format.space_after = Pt(3)
bullet.paragraph_format.line_spacing = 1.15

keyh = styles.add_style("Key Heading", WD_STYLE_TYPE.PARAGRAPH)
keyh.base_style = styles["Normal"]
keyh.paragraph_format.space_before = Pt(4)
keyh.paragraph_format.space_after = Pt(3)

sources_h = styles.add_style("Sources Heading", WD_STYLE_TYPE.PARAGRAPH)
sources_h.base_style = styles["Normal"]
sources_h.paragraph_format.space_before = Pt(5)
sources_h.paragraph_format.space_after = Pt(2)

sources_style = styles.add_style("Source List", WD_STYLE_TYPE.PARAGRAPH)
sources_style.base_style = styles["List Bullet"]
sources_style.paragraph_format.left_indent = Inches(0.25)
sources_style.paragraph_format.first_line_indent = Inches(-0.15)
sources_style.paragraph_format.space_after = Pt(1)
sources_style.paragraph_format.line_spacing = 1.0

# Quiet header and footer.
header_p = sec.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header_p.add_run("Maritime Insight | Maritime Economy (Sub-CPMK-4)"), size=8.5, color=MUTED)
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer_p.add_run("Materi inti presentasi kelompok | Halaman "), size=8.5, color=MUTED)
add_page_number(footer_p)

# Cover.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(72)
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("MATERI INTI PRESENTASI KELOMPOK"), size=11, color=MID_BLUE, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("Maritime Economy"), size=28, color=BLUE, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(32)
set_font(p.add_run("Sub-CPMK-4 | Mata Kuliah Maritime Insight (Wawasan Kemaritiman)"), size=13, color=MUTED)

table = doc.add_table(rows=3, cols=2)
table.autofit = False
table.columns[0].width = Inches(1.7)
table.columns[1].width = Inches(4.6)
for row in table.rows:
    row.cells[0].width = Inches(1.7)
    row.cells[1].width = Inches(4.6)
labels = [("Format", "8 subtopik; setiap subtopik ditujukan untuk 2 slide"), ("Bahasa", "Indonesia, ringkas dan akademik"), ("Catatan", "Data dan regulasi mengutamakan sumber resmi Indonesia")]
for row, (label, value) in zip(table.rows, labels):
    set_cell_shading(row.cells[0], "E8EEF5")
    for cell in row.cells:
        set_cell_margins(cell)
    pp = row.cells[0].paragraphs[0]
    set_font(pp.add_run(label), size=10, color=BLUE, bold=True)
    pp = row.cells[1].paragraphs[0]
    set_font(pp.add_run(value), size=10, color=INK)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(28)
set_font(p.add_run("Disusun untuk keperluan presentasi kelompok"), size=10, color=MUTED, italic=True)
doc.add_page_break()

for index, item in enumerate(sections):
    add_section(doc, item)
    if index < len(sections) - 1:
        doc.add_page_break()

doc.core_properties.title = "Materi Inti Maritime Economy - Sub-CPMK-4"
doc.core_properties.subject = "Maritime Insight (Wawasan Kemaritiman)"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
